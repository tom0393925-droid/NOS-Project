from pptx import Presentation

prs = Presentation('IM_Dashboard_User_Manual.pptx')
print(f"Slides: {len(prs.slides)}")
print(f"Width: {prs.slide_width.inches:.2f}\", Height: {prs.slide_height.inches:.2f}\"")
for i, slide in enumerate(prs.slides):
    texts = []
    for shape in slide.shapes:
        if shape.has_text_frame:
            t = shape.text_frame.text.strip()
            if t:
                texts.append(t[:70])
    preview = " | ".join(texts[:4])
    print(f"Slide {i+1:02d}: {preview}".encode('ascii', errors='replace').decode('ascii'))
